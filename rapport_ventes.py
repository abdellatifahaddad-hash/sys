#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Générateur de Rapport des Ventes - GUEDE OIL SARL
Script pour générer des rapports de ventes détaillés à partir du fichier MBACKE.xlsm

Usage:
    python3 rapport_ventes.py [--mois MOIS] [--destination DST] [--export FICHIER]
"""

import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from collections import defaultdict
import warnings
import sys
import argparse

warnings.filterwarnings('ignore')


class RapportVentes:
    """Classe pour générer les rapports de ventes"""
    
    def __init__(self, fichier_excel):
        """Initialise le générateur de rapport
        
        Args:
            fichier_excel: Chemin vers le fichier MBACKE.xlsm
        """
        self.fichier_excel = fichier_excel
        self.wb = None
        self.donnees_ventes = []
        self.donnees_factures = []
        
    def charger_donnees(self):
        """Charge les données depuis le fichier Excel"""
        print("📂 Chargement du fichier Excel...")
        try:
            self.wb = load_workbook(self.fichier_excel, read_only=True, data_only=True, keep_vba=False)
            print(f"✅ Fichier chargé avec succès: {len(self.wb.sheetnames)} feuilles trouvées")
            
            # Charger les données de ventes clients
            self._charger_etatclient()
            
            # Charger les données de factures
            self._charger_factures()
            
            return True
        except Exception as e:
            print(f"❌ Erreur lors du chargement: {e}")
            return False
    
    def _charger_etatclient(self):
        """Charge les données de la feuille etatclient (transactions clients)"""
        if 'etatclient' not in self.wb.sheetnames:
            print("⚠️  Feuille 'etatclient' non trouvée")
            return
        
        sheet = self.wb['etatclient']
        print("📊 Chargement des données de ventes clients...")
        
        # En-têtes attendus: DST, DATE, NOM PRENOM POMPISTE, CLIENT, IMMATRICULE, 
        #                    DESIGNATION, QUANTITE, PU VENTE, MTE, DSG, MOIS, REF
        
        ligne_debut = 2  # Les données commencent à la ligne 2
        count = 0
        
        for row in sheet.iter_rows(min_row=ligne_debut, values_only=True):
            if not row[0]:  # Ignorer les lignes vides
                continue
                
            try:
                vente = {
                    'destination': str(row[0]) if row[0] else '',
                    'date': row[1] if row[1] else None,
                    'pompiste': str(row[2]) if row[2] else '',
                    'client': str(row[3]) if row[3] else '',
                    'immatricule': str(row[4]) if row[4] else '',
                    'designation': str(row[5]) if row[5] else '',
                    'quantite': float(row[6]) if row[6] and str(row[6]) != '#N/A' else 0,
                    'prix_unitaire': float(row[7]) if row[7] and str(row[7]) != '#N/A' else 0,
                    'montant': float(row[8]) if row[8] and str(row[8]) != '#N/A' else 0,
                    'type': str(row[9]) if row[9] else '',
                    'mois': int(row[10]) if row[10] else 0,
                    'reference': str(row[11]) if row[11] else ''
                }
                self.donnees_ventes.append(vente)
                count += 1
            except Exception as e:
                # Ignorer les lignes avec erreurs
                pass
        
        print(f"✅ {count} transactions de ventes chargées")
    
    def _charger_factures(self):
        """Charge les données de la feuille FACTURES (factures fournisseurs)"""
        if 'FACTURES' not in self.wb.sheetnames:
            print("⚠️  Feuille 'FACTURES' non trouvée")
            return
        
        sheet = self.wb['FACTURES']
        print("📊 Chargement des données de factures...")
        
        # En-têtes: DST, TYP, Fournisseurs, Num Facture, Nature, DATE FACTURE, 
        #           MTE, MODE DE PAIMENT, MOIS
        
        ligne_debut = 3  # Les données commencent à la ligne 3
        count = 0
        
        for row in sheet.iter_rows(min_row=ligne_debut, values_only=True):
            if not row[0]:  # Ignorer les lignes vides
                continue
                
            try:
                facture = {
                    'destination': str(row[0]) if row[0] else '',
                    'type': str(row[1]) if row[1] else '',
                    'fournisseur': str(row[2]) if row[2] else '',
                    'numero': str(row[3]) if row[3] else '',
                    'nature': str(row[4]) if row[4] else '',
                    'date': row[5] if row[5] else None,
                    'montant': float(row[6]) if row[6] else 0,
                    'mode_paiement': str(row[7]) if row[7] else '',
                    'mois': int(row[8]) if row[8] else 0
                }
                self.donnees_factures.append(facture)
                count += 1
            except Exception as e:
                # Ignorer les lignes avec erreurs
                pass
        
        print(f"✅ {count} factures chargées")
    
    def generer_rapport_resume(self, mois=None, destination=None):
        """Génère un rapport résumé des ventes
        
        Args:
            mois: Numéro du mois à filtrer (1-12) ou None pour tous
            destination: Localisation à filtrer ou None pour toutes
            
        Returns:
            dict: Dictionnaire contenant les statistiques
        """
        donnees_filtrees = self._filtrer_donnees(self.donnees_ventes, mois, destination)
        
        if not donnees_filtrees:
            print("⚠️  Aucune donnée correspondante trouvée")
            return None
        
        # Calculs des statistiques
        stats = {
            'nombre_transactions': len(donnees_filtrees),
            'ventes_totales': sum(v['montant'] for v in donnees_filtrees),
            'par_destination': defaultdict(float),
            'par_type': defaultdict(float),
            'par_produit': defaultdict(lambda: {'quantite': 0, 'montant': 0}),
            'par_client': defaultdict(lambda: {'transactions': 0, 'montant': 0}),
            'par_mois': defaultdict(float),
        }
        
        # Agrégations
        for vente in donnees_filtrees:
            dst = vente['destination']
            typ = vente['type']
            produit = vente['designation']
            client = vente['client']
            mois_v = vente['mois']
            montant = vente['montant']
            quantite = vente['quantite']
            
            stats['par_destination'][dst] += montant
            stats['par_type'][typ] += montant
            stats['par_produit'][produit]['quantite'] += quantite
            stats['par_produit'][produit]['montant'] += montant
            stats['par_client'][client]['transactions'] += 1
            stats['par_client'][client]['montant'] += montant
            stats['par_mois'][mois_v] += montant
        
        return stats
    
    def _filtrer_donnees(self, donnees, mois=None, destination=None):
        """Filtre les données selon les critères
        
        Args:
            donnees: Liste des données à filtrer
            mois: Numéro du mois ou None
            destination: Destination ou None
            
        Returns:
            list: Données filtrées
        """
        resultat = donnees
        
        if mois is not None:
            resultat = [d for d in resultat if d.get('mois') == mois]
        
        if destination:
            resultat = [d for d in resultat if d.get('destination', '').upper() == destination.upper()]
        
        return resultat
    
    def afficher_rapport_console(self, mois=None, destination=None):
        """Affiche le rapport dans la console
        
        Args:
            mois: Numéro du mois à filtrer (1-12) ou None pour tous
            destination: Localisation à filtrer ou None pour toutes
        """
        print("\n" + "="*80)
        print("📊 RAPPORT DES VENTES - GUEDE OIL SARL".center(80))
        print("="*80 + "\n")
        
        # Filtres appliqués
        filtres = []
        if mois:
            mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                        'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
            filtres.append(f"Mois: {mois_noms[mois] if mois <= 12 else mois}")
        if destination:
            filtres.append(f"Destination: {destination}")
        
        if filtres:
            print("🔍 Filtres appliqués:", " | ".join(filtres))
            print()
        
        stats = self.generer_rapport_resume(mois, destination)
        
        if not stats:
            return
        
        # Vue d'ensemble
        print("📈 VUE D'ENSEMBLE")
        print("-" * 80)
        print(f"Nombre de transactions:     {stats['nombre_transactions']:>10}")
        print(f"Ventes totales:             {stats['ventes_totales']:>10,.0f} CFA")
        print()
        
        # Par destination
        if stats['par_destination']:
            print("🏢 VENTES PAR DESTINATION")
            print("-" * 80)
            for dst, montant in sorted(stats['par_destination'].items(), key=lambda x: x[1], reverse=True):
                pct = (montant / stats['ventes_totales'] * 100) if stats['ventes_totales'] > 0 else 0
                print(f"{dst:<30} {montant:>15,.0f} CFA  ({pct:>5.1f}%)")
            print()
        
        # Par type (CREDIT/ESPESE)
        if stats['par_type']:
            print("💳 VENTES PAR TYPE DE PAIEMENT")
            print("-" * 80)
            for typ, montant in sorted(stats['par_type'].items(), key=lambda x: x[1], reverse=True):
                pct = (montant / stats['ventes_totales'] * 100) if stats['ventes_totales'] > 0 else 0
                print(f"{typ:<30} {montant:>15,.0f} CFA  ({pct:>5.1f}%)")
            print()
        
        # Top 10 produits
        if stats['par_produit']:
            print("🛢️  TOP 10 PRODUITS")
            print("-" * 80)
            produits_tries = sorted(stats['par_produit'].items(), 
                                   key=lambda x: x[1]['montant'], reverse=True)[:10]
            for produit, data in produits_tries:
                print(f"{produit:<30} Qté: {data['quantite']:>10,.1f}  |  {data['montant']:>12,.0f} CFA")
            print()
        
        # Top 10 clients
        if stats['par_client']:
            print("👥 TOP 10 CLIENTS")
            print("-" * 80)
            clients_tries = sorted(stats['par_client'].items(), 
                                  key=lambda x: x[1]['montant'], reverse=True)[:10]
            for client, data in clients_tries:
                print(f"{client:<30} Trans: {data['transactions']:>5}  |  {data['montant']:>12,.0f} CFA")
            print()
        
        # Par mois
        if stats['par_mois'] and not mois:
            print("📅 VENTES PAR MOIS")
            print("-" * 80)
            mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                        'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
            for m in sorted(stats['par_mois'].keys()):
                montant = stats['par_mois'][m]
                nom_mois = mois_noms[m] if m <= 12 else f"Mois {m}"
                print(f"{nom_mois:<15} {montant:>15,.0f} CFA")
            print()
        
        print("="*80)
        print(f"📅 Généré le: {datetime.now().strftime('%d/%m/%Y à %H:%M:%S')}")
        print("="*80 + "\n")
    
    def exporter_excel(self, fichier_sortie, mois=None, destination=None):
        """Exporte le rapport vers un fichier Excel
        
        Args:
            fichier_sortie: Chemin du fichier Excel de sortie
            mois: Numéro du mois à filtrer ou None
            destination: Destination à filtrer ou None
        """
        print(f"📝 Export du rapport vers {fichier_sortie}...")
        
        stats = self.generer_rapport_resume(mois, destination)
        if not stats:
            return False
        
        # Créer un nouveau classeur
        wb_out = openpyxl.Workbook()
        wb_out.remove(wb_out.active)  # Supprimer la feuille par défaut
        
        # Style des en-têtes
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Feuille 1: Résumé
        ws_resume = wb_out.create_sheet("Résumé")
        ws_resume.append(["RAPPORT DES VENTES - GUEDE OIL SARL"])
        ws_resume.merge_cells('A1:C1')
        ws_resume['A1'].font = Font(size=16, bold=True)
        ws_resume['A1'].alignment = Alignment(horizontal='center')
        
        ws_resume.append([])
        ws_resume.append(["Généré le:", datetime.now().strftime('%d/%m/%Y à %H:%M:%S')])
        
        if mois or destination:
            ws_resume.append([])
            ws_resume.append(["FILTRES APPLIQUÉS"])
            if mois:
                mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                            'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
                ws_resume.append(["Mois:", mois_noms[mois] if mois <= 12 else str(mois)])
            if destination:
                ws_resume.append(["Destination:", destination])
        
        ws_resume.append([])
        ws_resume.append(["VUE D'ENSEMBLE"])
        ws_resume.append(["Nombre de transactions:", stats['nombre_transactions']])
        ws_resume.append(["Ventes totales (CFA):", stats['ventes_totales']])
        
        # Feuille 2: Par Destination
        if stats['par_destination']:
            ws_dest = wb_out.create_sheet("Par Destination")
            ws_dest.append(["Destination", "Montant (CFA)", "Pourcentage (%)"])
            for cell in ws_dest[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
            
            for dst, montant in sorted(stats['par_destination'].items(), key=lambda x: x[1], reverse=True):
                pct = (montant / stats['ventes_totales'] * 100) if stats['ventes_totales'] > 0 else 0
                ws_dest.append([dst, montant, pct])
            
            # Ajuster les largeurs de colonnes
            ws_dest.column_dimensions['A'].width = 30
            ws_dest.column_dimensions['B'].width = 20
            ws_dest.column_dimensions['C'].width = 15
        
        # Feuille 3: Par Produit
        if stats['par_produit']:
            ws_prod = wb_out.create_sheet("Par Produit")
            ws_prod.append(["Produit", "Quantité", "Montant (CFA)"])
            for cell in ws_prod[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
            
            produits_tries = sorted(stats['par_produit'].items(), 
                                   key=lambda x: x[1]['montant'], reverse=True)
            for produit, data in produits_tries:
                ws_prod.append([produit, data['quantite'], data['montant']])
            
            ws_prod.column_dimensions['A'].width = 35
            ws_prod.column_dimensions['B'].width = 15
            ws_prod.column_dimensions['C'].width = 20
        
        # Feuille 4: Par Client
        if stats['par_client']:
            ws_client = wb_out.create_sheet("Par Client")
            ws_client.append(["Client", "Nombre de transactions", "Montant (CFA)"])
            for cell in ws_client[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
            
            clients_tries = sorted(stats['par_client'].items(), 
                                  key=lambda x: x[1]['montant'], reverse=True)
            for client, data in clients_tries:
                ws_client.append([client, data['transactions'], data['montant']])
            
            ws_client.column_dimensions['A'].width = 35
            ws_client.column_dimensions['B'].width = 25
            ws_client.column_dimensions['C'].width = 20
        
        # Feuille 5: Détails des transactions
        ws_detail = wb_out.create_sheet("Détails Transactions")
        ws_detail.append(["Date", "Destination", "Client", "Produit", "Quantité", 
                         "Prix Unitaire", "Montant", "Type", "Pompiste"])
        for cell in ws_detail[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        donnees_filtrees = self._filtrer_donnees(self.donnees_ventes, mois, destination)
        for vente in sorted(donnees_filtrees, key=lambda x: x.get('date') or datetime(1900, 1, 1), reverse=True):
            date_str = vente['date'].strftime('%d/%m/%Y') if vente['date'] else ''
            ws_detail.append([
                date_str,
                vente['destination'],
                vente['client'],
                vente['designation'],
                vente['quantite'],
                vente['prix_unitaire'],
                vente['montant'],
                vente['type'],
                vente['pompiste']
            ])
        
        # Ajuster largeurs
        ws_detail.column_dimensions['A'].width = 12
        ws_detail.column_dimensions['B'].width = 25
        ws_detail.column_dimensions['C'].width = 30
        ws_detail.column_dimensions['D'].width = 20
        ws_detail.column_dimensions['E'].width = 12
        ws_detail.column_dimensions['F'].width = 15
        ws_detail.column_dimensions['G'].width = 15
        ws_detail.column_dimensions['H'].width = 12
        ws_detail.column_dimensions['I'].width = 25
        
        # Sauvegarder
        wb_out.save(fichier_sortie)
        print(f"✅ Rapport exporté avec succès!")
        return True
    
    def fermer(self):
        """Ferme le fichier Excel"""
        if self.wb:
            self.wb.close()


def main():
    """Fonction principale"""
    parser = argparse.ArgumentParser(
        description='Générateur de Rapport des Ventes - GUEDE OIL SARL',
        formatter_class=argparse.RawDescriptionHelpFormatter
    )
    
    parser.add_argument(
        '--fichier',
        default='sys/MBACKE.xlsm',
        help='Chemin vers le fichier Excel (défaut: sys/MBACKE.xlsm)'
    )
    
    parser.add_argument(
        '--mois',
        type=int,
        help='Filtrer par mois (1-12)'
    )
    
    parser.add_argument(
        '--destination',
        help='Filtrer par destination (ex: GUEDE-OIL-MBACKE1)'
    )
    
    parser.add_argument(
        '--export',
        help='Exporter vers un fichier Excel (ex: rapport_ventes.xlsx)'
    )
    
    args = parser.parse_args()
    
    # Créer le générateur de rapport
    rapport = RapportVentes(args.fichier)
    
    # Charger les données
    if not rapport.charger_donnees():
        sys.exit(1)
    
    # Afficher le rapport dans la console
    rapport.afficher_rapport_console(mois=args.mois, destination=args.destination)
    
    # Exporter si demandé
    if args.export:
        rapport.exporter_excel(args.export, mois=args.mois, destination=args.destination)
    
    # Fermer
    rapport.fermer()
    
    print("✅ Terminé!")


if __name__ == '__main__':
    main()
